"""PPTX parser: python-pptx text frames, tables and (plan P-5) speaker notes.
Locator `slide:<n>` and `slide:<n>#notes`. Embedded images are skipped (ADR-0014)."""

from __future__ import annotations

from app.services.ingestion import locators
from app.services.ingestion.errors import UnreadableDocumentError
from app.services.ingestion.parsers.base import ParsedDocument, Segment


def _shape_texts(shapes) -> list[str]:
    out: list[str] = []
    for shape in shapes:
        if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # GROUP
            out.extend(_shape_texts(shape.shapes))
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if text.strip():
                out.append(text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    out.append(" | ".join(cells))
    return out


def parse_pptx(path: str, *, include_notes: bool) -> ParsedDocument:
    from pptx import Presentation

    try:
        presentation = Presentation(path)
    except Exception as exc:  # noqa: BLE001 - corrupt/encrypted package
        raise UnreadableDocumentError("The PPTX file could not be read.") from exc

    segments: list[Segment] = []
    try:
        for index, slide in enumerate(presentation.slides, start=1):
            body = "\n".join(_shape_texts(slide.shapes))
            if body.strip():
                segments.append(Segment(body, locators.slide(index), page_no=index))
            if include_notes and slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    segments.append(Segment(notes, locators.slide(index, notes=True), page_no=index))
    except Exception as exc:  # noqa: BLE001
        raise UnreadableDocumentError("The PPTX file could not be read.") from exc
    return ParsedDocument(ext="pptx", segments=segments)
