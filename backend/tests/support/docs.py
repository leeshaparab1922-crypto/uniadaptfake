"""In-memory builders for PDF/PPTX/DOCX/TXT test documents (no external files)."""

from __future__ import annotations

import io


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_text_pdf(pages: list[list[str]]) -> bytes:
    """A real, minimal text PDF (Helvetica), one list of lines per page."""
    n = len(pages)
    page_ids = [4 + 2 * i for i in range(n)]
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {n} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, lines in enumerate(pages):
        content_id = 5 + 2 * i
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
                f"/Resources << /Font << /F1 3 0 R >> >> >>"
            ).encode()
        )
        stream = (
            "BT /F1 12 Tf 14 TL 50 740 Td " + " ".join(f"({_pdf_escape(ln)}) Tj T*" for ln in lines) + " ET"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for idx, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{idx} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_scanned_pdf(pages: list[str | None], *, size: tuple[int, int] = (1240, 1754)) -> bytes:
    """Image-only PDF (no text layer). `None` makes a blank white page."""
    from PIL import Image, ImageDraw, ImageFont

    images = []
    for text in pages:
        image = Image.new("RGB", size, "white")
        if text:
            font = ImageFont.load_default(size=48)
            draw = ImageDraw.Draw(image)
            y = 120
            for line in text.split("\n"):
                draw.text((100, y), line, fill="black", font=font)
                y += 80
        images.append(image)
    buf = io.BytesIO()
    images[0].save(buf, format="PDF", save_all=True, append_images=images[1:], resolution=150.0)
    return buf.getvalue()


def make_txt(text: str) -> bytes:
    return text.encode("utf-8")


def make_docx(paragraphs: list[tuple[str, str]]) -> bytes:
    """paragraphs: (style, text) e.g. ("Heading 1", "Unit 1") or ("Normal", "...")."""
    from docx import Document

    doc = Document()
    for style, text in paragraphs:
        if style.lower().startswith("heading"):
            doc.add_heading(text, level=int(style.split()[-1]))
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pptx(slides: list[tuple[str, str | None]]) -> bytes:
    """slides: (body text, speaker notes or None)."""
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[5]  # title only
    for body, notes in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = body
        if notes is not None:
            slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_legacy_ole() -> bytes:
    """Bytes that start with the OLE2 magic used by legacy .doc/.ppt files."""
    return bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 512
