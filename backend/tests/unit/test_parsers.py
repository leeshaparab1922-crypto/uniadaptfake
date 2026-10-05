"""ADR-0014 parsers: PDF (pdfplumber/pypdfium2), PPTX, DOCX, TXT. Locators page/slide/section/lines."""

from __future__ import annotations

import pytest

from app.services.ingestion.errors import UnreadableDocumentError
from app.services.ingestion.parsers.base import parse_document
from tests.support import docs


def _parse(tmp_path, name, data, ext, **kw):
    path = tmp_path / name
    path.write_bytes(data)
    defaults = dict(min_page_chars=25, render_dpi=72, include_pptx_notes=True)
    defaults.update(kw)
    return parse_document(str(path), ext, **defaults)


# ---------------------------------------------------------------- PDF


def test_pdf_text_pages_have_page_locators(tmp_path):
    long = "This is a page with plenty of selectable text content."
    doc = _parse(tmp_path, "a.pdf", docs.make_text_pdf([[long], [long + " two"]]), "pdf")
    assert [s.locator for s in doc.segments] == ["page:1", "page:2"]
    assert [s.page_no for s in doc.segments] == [1, 2]
    assert not any(s.needs_ocr for s in doc.segments)
    assert "plenty of selectable text" in doc.segments[0].text
    doc.close()


def test_pdf_scanned_pages_flagged_for_ocr_and_renderable(tmp_path):
    doc = _parse(tmp_path, "scan.pdf", docs.make_scanned_pdf(["Unit 1 scanned", None]), "pdf")
    assert [s.needs_ocr for s in doc.segments] == [True, True]
    png = doc.render_png(1)
    assert png.startswith(b"\x89PNG")
    doc.close()


def test_pdf_page_with_too_little_text_needs_ocr(tmp_path):
    doc = _parse(tmp_path, "a.pdf", docs.make_text_pdf([["7"]]), "pdf")
    assert doc.segments[0].needs_ocr
    doc.close()


def test_corrupt_file_unreadable_error_pdf(tmp_path):
    with pytest.raises(UnreadableDocumentError):
        _parse(tmp_path, "bad.pdf", b"%PDF-1.4\nnot really a pdf", "pdf")


# ---------------------------------------------------------------- PPTX


def test_pptx_slide_locators_and_notes(tmp_path):
    doc = _parse(
        tmp_path, "a.pptx", docs.make_pptx([("Intro to Trees", "Explain height"), ("Graphs", None)]), "pptx"
    )
    assert [(s.locator, s.page_no) for s in doc.segments] == [
        ("slide:1", 1),
        ("slide:1#notes", 1),
        ("slide:2", 2),
    ]
    assert doc.segments[1].text == "Explain height"


def test_pptx_notes_can_be_disabled(tmp_path):
    doc = _parse(
        tmp_path, "a.pptx", docs.make_pptx([("Trees", "secret notes")]), "pptx", include_pptx_notes=False
    )
    assert [s.locator for s in doc.segments] == ["slide:1"]


def test_pptx_slide_without_text_is_skipped_not_empty(tmp_path):
    doc = _parse(tmp_path, "a.pptx", docs.make_pptx([("", None), ("Real", None)]), "pptx")
    assert [s.locator for s in doc.segments] == ["slide:2"]


def test_corrupt_file_unreadable_error_pptx(tmp_path):
    with pytest.raises(UnreadableDocumentError):
        _parse(tmp_path, "bad.pptx", b"PK\x03\x04garbage", "pptx")


def test_pptx_embedded_image_not_extracted(tmp_path):
    doc = _parse(tmp_path, "a.pptx", docs.make_pptx([("Only text", None)]), "pptx")
    assert all(not s.needs_ocr for s in doc.segments)


# ---------------------------------------------------------------- DOCX


def test_docx_section_locator_uses_heading_path_and_paragraph_index(tmp_path):
    data = docs.make_docx(
        [
            ("Heading 1", "Unit 2"),
            ("Heading 2", "Trees"),
            ("Normal", "A tree is a graph."),
            ("Normal", "BST."),
        ]
    )
    doc = _parse(tmp_path, "a.docx", data, "docx")
    locs = [s.locator for s in doc.segments]
    assert locs[0] == "section:Unit 2#p1"
    assert locs[1] == "section:Unit 2 > Trees#p2"
    assert locs[2] == "section:Unit 2 > Trees#p3"
    assert doc.segments[2].text == "A tree is a graph."


def test_docx_heading_stack_pops_on_same_or_higher_level(tmp_path):
    data = docs.make_docx([("Heading 1", "A"), ("Heading 2", "A1"), ("Heading 1", "B"), ("Normal", "text")])
    doc = _parse(tmp_path, "a.docx", data, "docx")
    assert doc.segments[-1].locator == "section:B#p4"


def test_docx_blank_paragraphs_skipped(tmp_path):
    doc = _parse(tmp_path, "a.docx", docs.make_docx([("Normal", ""), ("Normal", "real")]), "docx")
    assert len(doc.segments) == 1


def test_corrupt_file_unreadable_error_docx(tmp_path):
    with pytest.raises(UnreadableDocumentError):
        _parse(tmp_path, "bad.docx", b"PK\x03\x04garbage", "docx")


def test_docx_without_headings_uses_document_locator(tmp_path):
    doc = _parse(tmp_path, "a.docx", docs.make_docx([("Normal", "hello")]), "docx")
    assert doc.segments[0].locator == "section:(document)#p1"


# ---------------------------------------------------------------- TXT


def test_txt_paragraph_line_ranges(tmp_path):
    text = "line one\nline two\n\n\nline five\n"
    doc = _parse(tmp_path, "a.txt", docs.make_txt(text), "txt")
    assert [(s.locator, s.text) for s in doc.segments] == [
        ("lines:1-2", "line one\nline two"),
        ("lines:5-5", "line five"),
    ]


def test_txt_strict_utf8_rejects_invalid(tmp_path):
    with pytest.raises(UnreadableDocumentError, match="UTF-8"):
        _parse(tmp_path, "a.txt", "café".encode("latin-1"), "txt")


def test_txt_whitespace_only_yields_no_segments(tmp_path):
    assert _parse(tmp_path, "a.txt", b"  \n\n\t\n", "txt").segments == []


def test_txt_bom_stripped(tmp_path):
    doc = _parse(tmp_path, "a.txt", b"\xef\xbb\xbfhello", "txt")
    assert doc.segments[0].text == "hello"


def test_txt_unicode_preserved(tmp_path):
    doc = _parse(tmp_path, "a.txt", "ಕನ್ನಡ ಪಠ್ಯ".encode(), "txt")
    assert doc.segments[0].text == "ಕನ್ನಡ ಪಠ್ಯ"


# ---------------------------------------------------------------- limits and tables (findings M5, m8)


def test_pdf_over_the_page_cap_is_unreadable_with_a_clear_message(tmp_path):
    data = docs.make_text_pdf([["page one has enough text to count"], ["page two"], ["page three"]])
    with pytest.raises(UnreadableDocumentError, match="3 pages; the limit is 2"):
        _parse(tmp_path, "big.pdf", data, "pdf", max_pdf_pages=2)


def test_pdf_at_the_page_cap_is_accepted(tmp_path):
    data = docs.make_text_pdf([["page one"], ["page two"]])
    doc = _parse(tmp_path, "ok.pdf", data, "pdf", max_pdf_pages=2)
    assert [s.locator for s in doc.segments] == ["page:1", "page:2"]


def test_docx_table_rows_become_one_segment_with_cells_joined(tmp_path):
    import io

    from docx import Document

    document = Document()
    document.add_heading("Unit 1", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Topic", "Hours"
    table.cell(1, 0).text, table.cell(1, 1).text = "Arrays", "4"
    buf = io.BytesIO()
    document.save(buf)
    doc = _parse(tmp_path, "t.docx", buf.getvalue(), "docx")
    table_text = [s.text for s in doc.segments if "|" in s.text]
    assert table_text == ["Topic | Hours\nArrays | 4"]


def test_pptx_table_and_grouped_shapes_are_extracted(tmp_path):
    import io

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    rows = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    rows.cell(0, 0).text, rows.cell(0, 1).text = "Algorithm", "Cost"
    rows.cell(1, 0).text, rows.cell(1, 1).text = "BFS", "O(V+E)"
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
    box.text_frame.text = "Grouped note about graphs"
    buf = io.BytesIO()
    prs.save(buf)
    doc = _parse(tmp_path, "t.pptx", buf.getvalue(), "pptx")
    text = doc.segments[0].text
    assert "Algorithm | Cost" in text and "BFS | O(V+E)" in text
    assert "Grouped note about graphs" in text
